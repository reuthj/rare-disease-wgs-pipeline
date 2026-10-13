#!/usr/bin/env python3
"""Stage 7: present Stage 6 candidate rows in a formatted Excel workbook."""
import math
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from main_classification import is_main_plp
import pandas as pd
from workflow_checks import validate_table, record_qc, is_mitochondrial_variant
import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SHEET_CANDIDATES = [
    ("Top_candidates", "top"),
    ("P_LP", "plp"),
    ("secondary", "secondary"),
    ("AD", "ad"),
    ("AR_homozygous", "ar_hom"),
    ("AR_compound", "ar_compound"),
    ("X_linked", "x_linked"),
    ("Mosaic", "mosaic"),
    ("Mitochondrial", "mitochondrial"),
    ("ALL", "all"),
]

COMMON_COLUMNS = [
    "Chr", "Pos", "REF_ALT", "rsID", "Gene_symbol", "Secondary_gene",
    "Transcript ID", "Transcript_biotype", "Consequences", "HGVS",
    "Gene_context", "Regulatory_context", "Protein_context", "genotype",
    "gnomAD_AF", "gnomAD counts", "inheritance_category", "ClinVar",
    "CADD", "Conservation", "LOEUF", "Missense_constraint",
    "Evidence_routes", "Route scores", "Nearest_splice_distance",
    "Nearest_TSS_distance", "Disease_associations", "Inheritance_and_mechanism",
    "ClinVar_pathogenic_variant_landscape", "Gene_HPO_terms", "Matched_HPO_terms",
    "Literature_support", "Top_publications", "Pathogenicity_evidence_coverage",
    "Variant_pathogenicity_score", "Inheritance_score", "Gene_disease_evidence_score",
    "Case_relevance_score", "Rarity_score", "Final_score",
]
COMPOUND_COLUMNS = ["Compound_pair_score", "Scored_compound_partner"]

# Compact, single-line report widths. Full cell values remain available in Excel.
WIDTHS = {
    "Chr": 7, "Pos": 12, "REF_ALT": 10, "rsID": 14, "Gene_symbol": 17,
    "Secondary_gene": 12, "Transcript ID": 22, "Transcript_biotype": 15,
    "Consequences": 21, "HGVS": 25, "Gene_context": 20,
    "Regulatory_context": 16, "Protein_context": 17, "genotype": 19,
    "gnomAD_AF": 13, "gnomAD counts": 14, "inheritance_category": 25,
    "ClinVar": 21, "CADD": 8, "Conservation": 16, "LOEUF": 9,
    "Missense_constraint": 13, "Evidence_routes": 16, "Route scores": 25,
    "Nearest_splice_distance": 22, "Nearest_TSS_distance": 22,
    "Disease_associations": 25, "Inheritance_and_mechanism": 29,
    "ClinVar_pathogenic_variant_landscape": 29, "Gene_HPO_terms": 24,
    "Matched_HPO_terms": 20, "Literature_support": 22, "Top_publications": 20,
    "Pathogenicity_evidence_coverage": 13, "Variant_pathogenicity_score": 12,
    "Inheritance_score": 12, "Gene_disease_evidence_score": 13,
    "Case_relevance_score": 12, "Rarity_score": 10, "Compound_pair_score": 12,
    "Scored_compound_partner": 22, "Final_score": 12,
}

REQUIRED = {
    "Chr_Pos", "Variant_REF_ALT", "Gene_symbol", "Gene_symbol_transcript_aligned",
    "Transcript_ID", "inheritance_category", "Case_gene_match", "Secondary_gene_match",
    "Secondary_MOI", "Secondary_reportable", "ClinVar_significance", "ClinVar_conditions",
    "CADD", "phyloP100way", "phastCons100way", "gnomAD_AF", "gnomAD_grpmax_AF",
    "gnomAD_het", "gnomAD_hom", "gnomAD_hemi", "Transcript_biotype", "Consequences",
    "HGVSc", "HGVSp", "Gene_context", "Regulatory_context", "GT", "DP", "AD_alt",
    "VAF", "Nearest_splice_distance", "Nearest_splice_transcript", "Nearest_TSS_distance",
    "Nearest_TSS_transcript", "Protein_context", "Missense_evidence", "Splice_prediction",
    "Regulatory_prediction", "LoF_evidence", "NMD_context", "Inframe_indel_evidence",
    "Evidence_routes", "Transcript_routes", "Known_variant", "ClinVar_pathogenic_fraction",
    "ClinVar_pathogenic_count", "Known_variant_landscape", "Disease_associations",
    "Missense_constraint", "LOEUF", "Gene_HPO_terms", "Matched_HPO_terms",
    "Inheritance_and_mechanism", "Literature_support", "Top_publications", "MANE_Transcript_ID",
    "Pathogenicity_evidence_coverage", "Variant_pathogenicity_score", "Inheritance_score",
    "Gene_disease_evidence_score", "Case_relevance_score", "Rarity_score", "Scored_gene",
    "Scored_transcript", "Scored_route", "Scored_compound_partner", "Compound_pair_score", "Final_score",
}

MISSING = {"", "na", "n/a", "nan", "none", "null", ".", "-"}
BLUE = "0563C1"
DARK = "17365D"
LIGHT_BLUE = "D9EAF7"
GRID = "D9E1F2"
BODY_FILL = "FCFCFC"
BODY_FONT = "Aptos"
BODY_FONT_SIZE = 11
HEADER_FONT_SIZE = 12
COLUMN_RULE = "A6A6A6"
HPO_COLUMN_RULE = "777777"


def is_missing(value: Any) -> bool:
    return value is None or str(value).strip().casefold() in MISSING


def clean(value: Any) -> str:
    return "" if is_missing(value) else str(value).strip()


def display_header(value: str) -> str:
    """Capitalize the first character of labels whose letters are all lowercase."""
    letters = value.replace("_", "").replace(" ", "")
    return value[0].upper() + value[1:] if value and letters.islower() else value


def split_pipe(value: Any) -> list[str]:
    text = clean(value).replace(r"\|", "|")
    return [part.strip() for part in re.split(r"\s*\|\s*", text)] if text else []


def split_list(value: Any) -> list[str]:
    text = clean(value)
    if not text:
        return []
    return list(dict.fromkeys(x.strip() for x in re.split(r"\s*(?:\|\||\||,|;)\s*", text) if x.strip()))


def truthy(value: Any) -> bool:
    return str(value).strip().casefold() in {"true", "1", "yes", "y", "t"}


def as_float(value: Any) -> float | None:
    try:
        if is_missing(value):
            return None
        result = float(str(value).replace(",", ""))
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def variant_key(row: pd.Series) -> str:
    return f"{clean(row.get('Chr_Pos'))}|{clean(row.get('Variant_REF_ALT'))}"


def transcript_map(row: pd.Series) -> dict[str, str]:
    txs = split_pipe(row.get("Transcript_ID"))

    raw_genes = row.get("Gene_symbol_transcript_aligned")
    raw_text = "" if raw_genes is None else str(raw_genes).strip()

    if not raw_text:
        genes = [""] * len(txs)
    else:
        genes = [
            "" if item.strip().casefold() in MISSING else item.strip()
            for item in re.split(r"\s*\|\s*", raw_text)
        ]
        # A lone "-" means no aligned gene was provided for any transcript.
        if len(genes) == 1 and not genes[0] and len(txs) > 1:
            genes = [""] * len(txs)

    if len(txs) != len(genes):
        raise ValueError(
            f"Transcript/gene alignment mismatch for {variant_key(row)}: "
            f"{len(txs)} transcripts, {len(genes)} genes."
        )

    return dict(zip(txs, genes))


def tx_blocks(value: Any) -> dict[str, str]:
    result: dict[str, str] = {}
    for block in split_pipe(value):
        if "=" in block:
            tx, text = block.split("=", 1)
            result[tx.strip()] = text.strip()
    return result


def aligned_tx_values(row: pd.Series, field: str, txs: list[str]) -> list[str]:
    """Return unique values for selected transcripts, removing accessions from HGVS."""
    blocks = tx_blocks(row.get(field))
    if blocks:
        values = [blocks[tx] for tx in txs if blocks.get(tx)]
    else:
        raw_values = split_pipe(row.get(field))
        if len(raw_values) == len(split_pipe(row.get("Transcript_ID"))):
            full_txs = split_pipe(row.get("Transcript_ID"))
            tx_to_gene = transcript_map(row)
            gene = clean(row.get("Scored_gene"))
            values = [
                value for tx, value in zip(full_txs, raw_values)
                if tx in txs and tx_to_gene.get(tx, "").casefold() == gene.casefold()
            ]
        else:
            values = raw_values
    cleaned = [compact_hgvs(value) for value in values]
    return list(dict.fromkeys(value for value in cleaned if value))


def gene_value(value: Any, gene: str, genes: list[str]) -> str:
    """Select a gene-prefixed block without splitting separators inside it."""
    text = clean(value)
    if not text:
        return ""
    tokens = sorted({g for g in genes if g}, key=len, reverse=True)
    if not tokens:
        return ""
    pattern = "|".join(re.escape(g) for g in tokens)
    marker = re.compile(rf"(?P<prefix>^|\s*\|\|\s*)(?P<gene>{pattern})\s*:\s*")
    hits = list(marker.finditer(text))
    if not hits:
        return text if len(set(tokens)) == 1 and gene in tokens else ""
    for i, hit in enumerate(hits):
        if hit.group("gene").casefold() == gene.casefold():
            end = hits[i + 1].start() if i + 1 < len(hits) else len(text)
            return re.sub(r"\s*\|\|\s*$", "", text[hit.end():end]).strip()
    return ""


def domains_only(value: str) -> str:
    domains: list[str] = []
    for item in re.split(r"\s*[,;]\s*", value):
        item = item.strip()
        if item.startswith("Domains="):
            item = item[len("Domains="):]
        if not item or item.startswith(("AA=", "Protein_pos=")):
            continue
        if item not in domains:
            domains.append(item)
    return "; ".join(domains)


def compact_protein(value: str) -> str:
    """Keep only a short list of recognized protein domains."""
    domain_text = ""
    for item in re.split(r"\s*;\s*", value):
        item = item.strip()
        if item.startswith("Domains="):
            domain_text = item.split("=", 1)[1].strip()

    kept_prefixes = ("Pfam:", "PANTHER:", "PROSITE_patterns:", "PROSITE_profiles:")
    domains = []
    for domain in domain_text.split(","):
        domain = domain.strip()
        if domain.startswith(kept_prefixes):
            short_name, _, identifier = domain.partition(":")
            label = short_name.replace("PROSITE_patterns", "PROSITE").replace("PROSITE_profiles", "PROSITE")
            compact = f"{label} {identifier}".strip()
            if compact and compact not in domains:
                domains.append(compact)

    return "; ".join(dict.fromkeys(domains))


def format_gene_context(value: str) -> str:
    """Capitalize all gene-context annotations and standardize exon spacing."""
    text = value.replace("\n", " | ")
    text = re.sub(r"\bexon\s*(\d+)\s*/\s*(\d+)\b", r"EXON \1/\2", text, flags=re.IGNORECASE)
    return text.upper()


def compact_hgvs(value: Any) -> str:
    """Remove transcript/protein accessions from a single HGVS value."""
    text = clean(value)
    text = re.sub(r"^(?:ENST|ENSP|NM_|NR_|NP_|XM_|XP_)[A-Za-z0-9_.-]*=", "", text)
    text = re.sub(r"^(?:ENST|ENSP|NM_|NR_|NP_|XM_|XP_)[A-Za-z0-9_.-]*:", "", text)
    return text.strip()


def cadd_phred(value: Any) -> float | None:
    text = clean(value)
    match = re.search(r"(?:^|[;|,]\s*)PHRED\s*=\s*([-+]?(?:\d+(?:\.\d*)?|\.\d+))", text, re.IGNORECASE)
    if match:
        return as_float(match.group(1))
    return as_float(text)


def compact_nmd(value: Any) -> str:
    text = clean(value)
    if not text:
        return ""
    match = re.search(r"NMD(?:_context)?\s*=\s*([^;,]+)", text, re.IGNORECASE)
    if match:
        text = match.group(1).strip()
    text = re.sub(r"^NMD_", "", text, flags=re.IGNORECASE)
    text = re.sub(r"_variant$", "", text, flags=re.IGNORECASE)
    text = text.replace("_", " ").strip()
    return f"NMD={text}" if text else ""


def compact_lof(value: Any) -> str:
    """Keep the short LoFTEE classification and decisive rule flags."""
    text = clean(value)
    if not text:
        return ""
    wanted = {"LOFTEE", "FILTER", "50_BP_RULE"}
    found: list[str] = []
    for token in re.split(r"[;,]", text):
        token = token.strip()
        if "=" not in token:
            continue
        key, val = (part.strip() for part in token.split("=", 1))
        key_upper = key.upper()
        if key_upper not in wanted or not val:
            continue
        if key_upper == "50_BP_RULE":
            item = f"50bp={val}"
        else:
            item = f"{key_upper}={val}"
        if item and item not in found:
            found.append(item)
    return "; ".join(found)


def number_or_original(value: Any) -> float | int | str:
    number = as_float(value)
    if number is None:
        return clean(value)
    return int(number) if number.is_integer() else number


def route_map(value: Any, txs: list[str]) -> dict[str, list[str]]:
    out = {tx: [] for tx in txs}
    for block in split_pipe(value):
        if ":" not in block:
            continue
        raw_tx, raw_routes = block.split(":", 1)
        match = next((tx for tx in txs if tx.split(".", 1)[0] == raw_tx.strip().split(".", 1)[0]), None)
        if match:
            out[match] = list(dict.fromkeys(r.strip().upper() for r in raw_routes.split(",") if r.strip()))
    return out


def rich(parts: list[tuple[str, dict[str, Any] | None]]) -> CellRichText | str:
    one_line_parts = [
        (re.sub(r"\s*[\r\n]+\s*", " | ", text), style)
        for text, style in parts
        if text
    ]
    blocks = [TextBlock(InlineFont(**style), text) if style else text for text, style in one_line_parts]
    if not blocks:
        return ""
    return CellRichText(*blocks)


def gene_and_transcripts(row: pd.Series, mane_ids: set[str]) -> tuple[str, str, list[str], dict[str, str]]:
    gene = clean(row.get("Scored_gene"))
    mapping = transcript_map(row)
    selected = [tx for tx, tx_gene in mapping.items() if tx_gene.casefold() == gene.casefold()]
    if gene and not selected and mapping and all(not tx_gene for tx_gene in mapping.values()):
        scored_tx = clean(row.get("Scored_transcript"))
        matched_tx = next(
            (tx for tx in mapping if tx == scored_tx or tx.split(".", 1)[0] == scored_tx.split(".", 1)[0]),
            None,
        )
        if matched_tx:
            selected = [matched_tx]
    if gene and not selected:
        # Never display another gene's transcript data for a scored hypothesis.
        raise ValueError(f"Scored_gene {gene!r} has no aligned transcript for {variant_key(row)}")
    return gene, "", selected, mapping


def format_routes(row: pd.Series, scored_txs: list[str]) -> tuple[CellRichText | str, str]:
    mapping = transcript_map(row)
    all_routes = route_map(row.get("Transcript_routes"), list(mapping))
    gene = clean(row.get("Scored_gene"))
    relevant_txs = [tx for tx, tx_gene in mapping.items() if tx_gene.casefold() == gene.casefold()]
    if not relevant_txs and gene and all(not tx_gene for tx_gene in mapping.values()):
        relevant_txs = [tx for tx in scored_txs if tx in mapping]
    selected_route = clean(row.get("Scored_route")).upper()
    if gene and is_missing(row.get("Scored_route")):
        raise ValueError(
            "Stage 7 needs Scored_route in the Stage 6 table. Expose "
            "candidate.pathogenicity_route from Stage 6 before running Stage 7."
        )
    if selected_route == "COMMON":
        selected_route = ""
    
    bold_routes = {selected_route} if selected_route else set()
    scored_first = [selected_route] if selected_route else []
    other_routes = [r for tx in relevant_txs for r in all_routes.get(tx, []) if r != "COMMON" and r != selected_route]
    routed = list(dict.fromkeys(scored_first + other_routes))
    score_routes = list(routed)
    route_parts: list[tuple[str, dict[str, Any] | None]] = []
    for i, route in enumerate(routed):
        if i:
            route_parts.append((" | ", None))
        route_parts.append((route, {"b": True} if route in bold_routes else None))
    route_text = rich(route_parts)
    route_score_parts: list[tuple[str, dict[str, Any] | None]] = []
    source_fields = {
        "MISSENSE": ["Missense_evidence"],
        "SPLICE": ["Splice_prediction", "Splice_context"],
        "REGULATORY": ["Regulatory_context"],
        "LOF": ["LoF_evidence", "NMD_context"],
        "INFRAME": ["Inframe_indel_evidence"],
        "INFRAME_INDEL": ["Inframe_indel_evidence"],
        "IN-FRAME": ["Inframe_indel_evidence"],
    }
    for route in score_routes:
        components: list[str] = []
        if route == "LOF":
            lof_blocks = tx_blocks(row.get("LoF_evidence"))
            nmd_blocks = tx_blocks(row.get("NMD_context"))
            for tx in relevant_txs:
                lof_text = compact_lof(lof_blocks.get(tx))
                if lof_text:
                    components.append(lof_text)
                nmd_text = compact_nmd(nmd_blocks.get(tx)) or compact_nmd(lof_blocks.get(tx))
                if nmd_text:
                    components.append(nmd_text)
            components = list(dict.fromkeys(components))
            if components:
                if len(score_routes) > 1:
                    route_score_parts.append((route + ": ", {"b": True}))
                route_score_parts.append(("; ".join(components), None))
            continue
        for field in source_fields.get(route, []):
            blocks = tx_blocks(row.get(field))
            for tx in relevant_txs:
                if blocks.get(tx):
                    label = field.replace("_evidence", "").replace("_prediction", "")
                    if field == "Splice_context":
                        label = "context"
                    components.append(f"{label}={blocks[tx]}")
        if (
            route == "REGULATORY"
            and clean(row.get("Regulatory_prediction"))
            and not any("remm" in component.casefold() for component in components)
        ):
            remm_value = re.sub(
                r"^\s*ReMM\s*=\s*", "", clean(row.get("Regulatory_prediction")), flags=re.IGNORECASE
            )
            components.append(f"ReMM={remm_value}")
        components = list(dict.fromkeys(
            re.sub(r"\bReMM\s*=\s*ReMM\s*=", "ReMM=", component, flags=re.IGNORECASE)
            for component in components
        ))
        if components:
            prefix = f"{route}: " if len(score_routes) > 1 else ""
            if prefix:
                route_score_parts.append((route + ": ", {"b": True}))
            route_score_parts.append(("; ".join(dict.fromkeys(components)), None))
    combined: list[tuple[str, dict[str, Any] | None]] = []
    for i, part in enumerate(route_score_parts):
        if i and part[1] is not None:
            combined.append((" | ", None))
        combined.append(part)
    return route_text, rich(combined)


def display_row(row: pd.Series, case_genes: set[str]) -> tuple[dict[str, Any], set[str]]:
    gene, _, txs, mapping = gene_and_transcripts(row, set())
    all_genes = list(dict.fromkeys(x for x in mapping.values() if x and x != "-"))
    if not all_genes and gene:
        all_genes = [gene]
    mane = set(split_pipe(row.get("MANE_Transcript_ID")))
    txs_mane = [tx for tx in txs if tx in mane or tx.split(".", 1)[0] in {x.split(".", 1)[0] for x in mane}]
    scored_tx = clean(row.get("Scored_transcript"))
    scored_tx_full = next((tx for tx in txs if tx == scored_tx), None)
    if scored_tx_full is None and scored_tx:
        scored_tx_full = next((tx for tx in txs if tx.split(".", 1)[0] == scored_tx.split(".", 1)[0]), None)
    scored_txs = [scored_tx_full] if scored_tx_full else []

    def aligned(field: str) -> list[str]:
        vals = split_pipe(row.get(field))
        source_txs = list(mapping)
        if len(vals) == len(source_txs):
            if gene and not any(mapping.values()):
                return [val for tx, val in zip(source_txs, vals) if tx in txs]
            return [val for tx, val in zip(source_txs, vals) if mapping[tx].casefold() == gene.casefold()]
        return [clean(row.get(field))] if clean(row.get(field)) else []

    tx_parts = []

    for i, tx in enumerate(txs):
        if tx in txs_mane and tx in scored_txs:
            style = {"b": True, "color": BLUE}
        elif tx in txs_mane:
            style = {"b": True}
        elif tx in scored_txs:
            style = {"color": BLUE}
        else:
            style = None

        tx_parts.append((tx, style))

        if i < len(txs) - 1:
            tx_parts.append((" | ", None))

    tx_display = rich(tx_parts)

    # Gene-level values are parsed only when explicitly gene-prefixed; a single-gene
    # unprefixed value is safe to retain as-is.
    genes = all_genes
    missense_constraint = gene_value(row.get("Missense_constraint"), gene, genes)
    loeuf_value = gene_value(row.get("LOEUF"), gene, genes)
    clinvar_fraction = gene_value(row.get("ClinVar_pathogenic_fraction"), gene, genes)
    clinvar_count = gene_value(row.get("ClinVar_pathogenic_count"), gene, genes)
    disease = gene_value(row.get("Disease_associations"), gene, genes)
    inheritance_mechanism = gene_value(row.get("Inheritance_and_mechanism"), gene, genes)
    hpo = gene_value(row.get("Gene_HPO_terms"), gene, genes)
    matched_hpo = gene_value(row.get("Matched_HPO_terms"), gene, genes)
    literature = gene_value(row.get("Literature_support"), gene, genes)
    top_publications = gene_value(row.get("Top_publications"), gene, genes)
    known_landscape = gene_value(row.get("Known_variant_landscape"), gene, genes)

    # Transcript-associated fields are reduced to the selected Scored_gene transcripts.
    gene_context_blocks = tx_blocks(row.get("Gene_context"))
    gene_context = format_gene_context(
        "\n".join(gene_context_blocks[tx] for tx in txs if gene_context_blocks.get(tx))
    )
    protein_blocks = tx_blocks(row.get("Protein_context"))
    protein = "; ".join(dict.fromkeys(compact_protein(protein_blocks[tx]) for tx in txs if protein_blocks.get(tx)))
    regulatory = clean(row.get("Regulatory_context"))
    regulatory = regulatory.replace(" ; ", " | ").replace("; ", " | ")

    routes, route_scores = format_routes(row, scored_txs)
    mitochondrial = is_mitochondrial_variant(row)
    if mitochondrial:
        route_scores += " | Partial MT evaluation; MT-specific pathogenicity and rarity unavailable"
    clinvar_class = clean(row.get("ClinVar_significance"))
    clinvar_conditions = clean(row.get("ClinVar_conditions"))
    clinvar = " | ".join(x for x in [clinvar_class, clinvar_conditions] if x)
    if clinvar_conditions and clinvar_class and clinvar_class.casefold() in clinvar_conditions.casefold():
        clinvar = clinvar_conditions

    chr_pos = clean(row.get("Chr_Pos"))
    chrom, _, pos = chr_pos.partition(":")
    splice_dist = as_float(row.get("Nearest_splice_distance"))
    splice_tx = clean(row.get("Nearest_splice_transcript"))

    if (
        splice_dist is None
        or abs(splice_dist) >= 50
        or (
            splice_tx
            and not any(
                splice_tx.split(".", 1)[0] == tx.split(".", 1)[0]
                for tx in txs
            )
        )
    ):
        splice_display = ""
    else:
        splice_display = " | ".join(
            x for x in [
                clean(row.get("Nearest_splice_distance")),
                splice_tx,
            ] if x
        )

    tss_display = " | ".join(
        x for x in [
            clean(row.get("Nearest_TSS_distance")),
            clean(row.get("Nearest_TSS_transcript")),
        ] if x
    )

    # Case status is determined from the selected hypothesis gene against case_context.
    is_case = gene.casefold() in {x.casefold() for x in case_genes}
    secondary = ""
    if truthy(row.get("Secondary_gene_match")) or truthy(row.get("Secondary_reportable")):
        secondary = "TRUE"
        moi = clean(row.get("Secondary_MOI"))
        if moi:
            secondary += f" | MOI: {moi}"
        if disease:
            secondary += f" | {disease}"

    frac = clean(clinvar_fraction)
    count = clean(clinvar_count)
    landscape = clean(known_landscape)
    landscape_parts = [f"fraction={frac}" if frac else "", f"count={count}" if count else "", landscape]
    clinvar_landscape = " | ".join(x for x in landscape_parts if x)

    gt = clean(row.get("GT"))
    ad_dp = "/".join(x for x in [clean(row.get("AD_alt")), clean(row.get("DP"))] if x)
    vaf = as_float(row.get("VAF"))
    genotype = " | ".join(x for x in [gt, ad_dp, f"{vaf:.3f}" if vaf is not None else ""] if x)
    count_fields = ["gnomAD_het", "gnomAD_hom"]
    if chrom.casefold().removeprefix("chr") == "x":
        count_fields.append("gnomAD_hemi")
    count_values = [clean(row.get(field)) for field in count_fields]
    gnomad_counts = " | ".join(count_values) if any(count_values) else ""
    conservation = " | ".join(x for x in [f"phyloP: {clean(row.get('phyloP100way'))}" if clean(row.get("phyloP100way")) else "", f"phastCons: {clean(row.get('phastCons100way'))}" if clean(row.get("phastCons100way")) else ""] if x)
    gnomad_af = as_float(row.get("gnomAD_AF"))
    if mitochondrial:
        gnomad_af = as_float(row.get("gnomAD_mt_AF"))
        gnomad_counts = " | ".join(
            f"{label}={clean(row.get(field))}"
            for label, field in [("homoplasmic_AF", "gnomAD_mt_AF_hom"), ("heteroplasmic_AF", "gnomAD_mt_AF_het")]
            if clean(row.get(field))
        )
    elif gnomad_af is None:
        gnomad_af = 0.0
    coverage = as_float(row.get("Pathogenicity_evidence_coverage"))
    if coverage is not None and coverage.is_integer():
        coverage = int(coverage)

    hgvs_c = aligned_tx_values(row, "HGVSc", txs)
    hgvs_p = aligned_tx_values(row, "HGVSp", txs)
    hgvs_values = []
    for i in range(max(len(hgvs_c), len(hgvs_p))):
        pair = [
            values[i] for values in (hgvs_c, hgvs_p)
            if i < len(values) and values[i]
        ]
        value = ", ".join(pair)
        if value and value not in hgvs_values:
            hgvs_values.append(value)

    data = {
        "Chr": chrom.removeprefix("chr"),
        "Pos": int(pos) if pos.isdigit() else pos,
        "REF_ALT": clean(row.get("Variant_REF_ALT")),
        "rsID": "" if clean(row.get("Known_variant")).casefold() == "norsid" else clean(row.get("Known_variant")),
        "Gene_symbol": gene,
        "Secondary_gene": secondary,
        "Transcript ID": tx_display,
        "Transcript_biotype": " | ".join(aligned("Transcript_biotype")),
        "Consequences": " | ".join(aligned("Consequences")),
        "HGVS": " | ".join(hgvs_values),
        "Gene_context": gene_context,
        "Regulatory_context": regulatory,
        "Protein_context": protein,
        "genotype": genotype,
        "gnomAD_AF": gnomad_af,
        "gnomAD counts": gnomad_counts,
        "inheritance_category": clean(row.get("inheritance_category")),
        "ClinVar": rich([(clinvar_class, {"b": True}), (" | " if clinvar_class and clinvar_conditions else "", None), (clinvar_conditions, None)]) if clinvar else "",
        "CADD": cadd_phred(row.get("CADD")),
        "Conservation": conservation,
        "LOEUF": number_or_original(loeuf_value),
        "Missense_constraint": missense_constraint,
        "Evidence_routes": routes,
        "Route scores": route_scores,
        "Nearest_splice_distance": splice_display,
        "Nearest_TSS_distance": tss_display,
        "Disease_associations": disease,
        "Inheritance_and_mechanism": inheritance_mechanism,
        "ClinVar_pathogenic_variant_landscape": clinvar_landscape,
        "Gene_HPO_terms": hpo,
        "Matched_HPO_terms": matched_hpo,
        "Literature_support": literature,
        "Top_publications": top_publications,
        "Pathogenicity_evidence_coverage": coverage,
        "Variant_pathogenicity_score": as_float(row.get("Variant_pathogenicity_score")),
        "Inheritance_score": as_float(row.get("Inheritance_score")),
        "Gene_disease_evidence_score": as_float(row.get("Gene_disease_evidence_score")),
        "Case_relevance_score": as_float(row.get("Case_relevance_score")),
        "Rarity_score": as_float(row.get("Rarity_score")),
        "Compound_pair_score": as_float(row.get("Compound_pair_score")),
        "Scored_compound_partner": clean(row.get("Scored_compound_partner")),
        "Final_score": as_float(row.get("Final_score")),
    }
    return data, {"is_case": is_case, "mane": txs_mane, "scored_txs": scored_txs, "clinvar_class": clinvar_class}


def expected_inheritance(context: dict[str, Any]) -> set[str]:
    aliases = {"AD_singleton": "AD", "X_linked_hemi": "X_linked", "X_linked_hemizygous": "X_linked", "mtDNA": "Mitochondrial"}
    out = {aliases.get(str(x).strip(), str(x).strip()) for x in (context.get("expected_inheritance") or [])}
    if "AR" in out:
        out.remove("AR")
        out.update({"AR_homozygous", "AR_compound_heterozygous"})
    return out


def row_inheritance(row: pd.Series) -> set[str]:
    aliases = {"AD_singleton": "AD", "X_linked_hemi": "X_linked", "X_linked_hemizygous": "X_linked", "mtDNA": "Mitochondrial"}
    return {aliases.get(x, x) for x in split_list(row.get("inheritance_category"))}


def valid_compound_variant_indices(frame: pd.DataFrame) -> set[int]:
    by_gene_and_variant = {
        (clean(row.get("Scored_gene")), variant_key(row)): idx
        for idx, row in frame.iterrows()
        if clean(row.get("Scored_gene"))
    }
    paired_indices: set[int] = set()
    for idx, row in frame.iterrows():
        gene = clean(row.get("Scored_gene"))
        if not gene or "AR_compound_heterozygous" not in row_inheritance(row):
            continue
        score = as_float(row.get("Compound_pair_score"))
        partner_key = clean(row.get("Scored_compound_partner"))
        if score is None or score < 30 or not partner_key:
            continue
        partner_idx = by_gene_and_variant.get((gene, partner_key))
        if partner_idx is None or partner_idx == idx:
            continue
        partner_row = frame.loc[partner_idx]
        if "AR_compound_heterozygous" not in row_inheritance(partner_row):
            continue
        partner_score = as_float(partner_row.get("Compound_pair_score"))
        if partner_score is None or partner_score < 30:
            continue
        if clean(partner_row.get("Scored_compound_partner")) != variant_key(row):
            continue
        paired_indices.update({idx, partner_idx})
    return paired_indices


def apply_compound_singleton_fallback(frame: pd.DataFrame) -> pd.DataFrame:
    """Use AD only for a selected alternative gene; preserve same-gene AR."""
    validate_table(frame, "Stage 7 input", required=["Scored_inheritance", "AD_fallback", "MANE_Transcript_ID"], aligned=True)
    result = frame.copy()
    paired_indices = valid_compound_variant_indices(result)
    for idx, row in result.iterrows():
        categories = split_list(row.get("inheritance_category"))
        if "AR_compound_heterozygous" not in categories:
            continue
        if idx in paired_indices:
            continue
        if not truthy(row.get("AD_fallback")):
            if clean(row.get("Scored_inheritance")) == "AR_compound_heterozygous":
                continue
            if not clean(row.get("Scored_gene")) and not clean(row.get("Scored_inheritance")):
                continue
            raise ValueError("Alternative-gene AD fallback requires recalculated Stage 6 AD scores")
        if clean(row.get("Scored_inheritance")) != "AD":
            raise ValueError("AD fallback has inconsistent scored inheritance")
        if clean(row.get("Scored_gene")).casefold() in {gene.casefold() for gene in split_list(row.get("Compound_het_gene"))}:
            raise ValueError("AD fallback cannot use the same AR compound gene")
        if not is_missing(row.get("Scored_compound_partner")) or not is_missing(row.get("Compound_pair_score")):
            raise ValueError("AD fallback retains selected compound fields")
        categories = [category for category in categories if category != "AR_compound_heterozygous"]
        if "AD" not in categories:
            categories.append("AD")
        result.at[idx, "inheritance_category"] = " | ".join(categories)
    return result


def make_tab_frames(frame: pd.DataFrame, views: list[dict[str, Any]], expected: set[str]) -> dict[str, list[int]]:
    scores = frame["Final_score"].map(as_float)
    case = pd.Series([v["is_case"] for v in views], index=frame.index)
    plp = frame["ClinVar_significance"].map(is_main_plp)
    categories = frame["inheritance_category"].map(lambda x: row_inheritance(pd.Series({"inheritance_category": x})))
    compound_flag = frame["Scored_compound_partner"].map(lambda x: not is_missing(x)) & frame["Compound_pair_score"].map(
        lambda x: (as_float(x) if as_float(x) is not None else -1) >= 30
    )
    paired_indices = valid_compound_variant_indices(frame)
    x_chrom = frame["Chr_Pos"].map(lambda v: clean(v).split(":", 1)[0].casefold().removeprefix("chr") == "x")
    qualifies = scores.ge(30)
    selected: dict[str, list[int]] = {name: [] for name, _ in SHEET_CANDIDATES}
    selected["ALL"] = list(frame.index)
    selected["P_LP"] = list(frame.index[plp])
    selected["secondary"] = list(frame.index[frame["Secondary_reportable"].map(truthy)])
    selected["AD"] = list(frame.index[qualifies & categories.map(lambda x: "AD" in x)])
    selected["AR_homozygous"] = list(frame.index[qualifies & categories.map(lambda x: "AR_homozygous" in x)])
    selected["AR_compound"] = [i for i in frame.index if i in paired_indices]
    selected["X_linked"] = list(frame.index[qualifies & x_chrom])
    selected["Mosaic"] = list(frame.index[qualifies & categories.map(lambda x: "Mosaic" in x)])
    selected["Mitochondrial"] = list(frame.index[qualifies & categories.map(lambda x: "Mitochondrial" in x)])

    expected_rows = frame.index[categories.map(lambda values: bool(values & expected))]
    case_idx = list(frame.index[case & qualifies])
    pathogenicity_qualifies = frame["Variant_pathogenicity_score"].map(as_float).ge(30)
    # Each non-case allele must qualify on its own pathogenicity, including
    # alleles ranked using a strong compound-pair score.
    noncase_idx = [i for i in expected_rows if not case.loc[i] and qualifies.loc[i]
                   and pathogenicity_qualifies.loc[i]]
    # Preserve every case-gene row; target max(case rows + 10, 20) candidates.
    noncase_limit = max(10, 20 - len(case_idx))
    top_noncase = sorted(
        noncase_idx,
        key=lambda i: (-((as_float(frame.loc[i, "Compound_pair_score"]) if compound_flag.loc[i] else None) or scores.loc[i] or 0), i),
    )[:noncase_limit]
    selected["Top_candidates"] = list(dict.fromkeys(case_idx + top_noncase))
    return selected


def sort_indices(frame: pd.DataFrame, indices: list[int], view_by_index: dict[int, dict[str, Any]], compound_tab: bool = False) -> list[int]:
    order = {idx: pos for pos, idx in enumerate(frame.index)}
    genes: dict[str, list[int]] = defaultdict(list)
    for idx in indices:
        genes[clean(frame.loc[idx, "Scored_gene"])].append(idx)
    case_flag = lambda idx: bool(view_by_index[idx]["is_case"])
    def pair_rank(idx: int) -> float:
        if is_missing(frame.loc[idx, "Scored_compound_partner"]):
            return -1.0
        return as_float(frame.loc[idx, "Compound_pair_score"]) or -1.0
    def row_score(idx: int) -> float:
        return as_float(frame.loc[idx, "Final_score"]) or -1.0
    gene_rank = {}
    for gene, idxs in genes.items():
        max_pair = max((pair_rank(i) for i in idxs), default=-1.0)
        max_score = max((row_score(i) for i in idxs), default=-1.0)
        gene_rank[gene] = (not any(case_flag(i) for i in idxs), -(max_pair if max_pair >= 30 else max_score), gene.casefold())
    ordered_genes = sorted(genes, key=lambda g: gene_rank[g])
    output: list[int] = []
    for gene in ordered_genes:
        idxs = genes[gene]
        paired = [i for i in idxs if pair_rank(i) >= 30]
        paired.sort(key=lambda i: (-pair_rank(i), order[i]))
        paired_set = set(paired)
        rest = [i for i in idxs if i not in paired_set]
        rest.sort(key=lambda i: (-row_score(i), order[i]))
        output.extend(paired + rest)
    return output


def rich_gene(gene: str, is_case: bool) -> CellRichText | str:
    parts = [(gene, {"color": BLUE})]
    if is_case:
        parts.append((" | CASE", {"b": True, "color": DARK}))
    return rich(parts)


def load_inputs(variants_path: Path, context_path: Path) -> tuple[pd.DataFrame, dict[str, Any]]:
    if not variants_path.is_file():
        raise FileNotFoundError(f"Stage 6 input table not found: {variants_path}")
    if not context_path.is_file():
        raise FileNotFoundError(f"case_context not found: {context_path}")
    frame = pd.read_csv(variants_path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    if "Scored_route" not in frame.columns:
        raise ValueError(
            "Stage 6 table does not contain Scored_route. Add candidate.pathogenicity_route "
            "to Stage 6 output; Stage 7 will not guess the winning route."
        )
    missing = sorted(REQUIRED - set(frame.columns))
    if missing:
        raise ValueError("Stage 6 table is missing required columns: " + ", ".join(missing))
    if frame.empty:
        raise ValueError("Stage 6 table contains no rows.")
    with context_path.open(encoding="utf-8") as handle:
        context = yaml.safe_load(handle) or {}
    if not isinstance(context, dict):
        raise ValueError("case_context.yaml must contain a YAML mapping/object at its root.")
    return frame.reset_index(drop=True), context


def write_workbook(frame: pd.DataFrame, context: dict[str, Any], output: Path) -> None:
    print("[Stage 7] Preparing report rows...", flush=True)
    frame = apply_compound_singleton_fallback(frame)
    case_genes = {str(x).strip() for x in (context.get("case_genes") or []) if str(x).strip()}
    expected = expected_inheritance(context)
    views: dict[int, dict[str, Any]] = {}
    display: dict[int, dict[str, Any]] = {}
    total = len(frame)
    progress_step = max(100, total // 10)
    for processed, (idx, row) in enumerate(frame.iterrows(), start=1):
        display[idx], views[idx] = display_row(row, case_genes)
        views[idx]["case_gene"] = views[idx]["is_case"]
        if processed == 1 or processed % progress_step == 0 or processed == total:
            print(f"[Stage 7] Formatted {processed:,}/{total:,} variants.", flush=True)
    tabs = make_tab_frames(frame, [views[i] for i in frame.index], expected)

    print("[Stage 7] Building workbook tabs...", flush=True)
    wb = Workbook()
    wb.remove(wb.active)
    headers = COMMON_COLUMNS
    for sheet_name, _ in SHEET_CANDIDATES:
        indices = tabs.get(sheet_name, [])
        print(f"[Stage 7] {sheet_name}: {len(indices):,} variants.", flush=True)
        if sheet_name != "ALL" and not indices:
            continue
        compound = sheet_name in {"Top_candidates", "P_LP", "secondary", "AR_compound", "ALL"}
        cols = (headers[:-1] + COMPOUND_COLUMNS + headers[-1:]) if compound else headers
        ws = wb.create_sheet(sheet_name)
        ws.append(cols)
        ordered = sort_indices(frame, indices, views, compound_tab=compound)
        for idx in ordered:
            row = dict(display[idx])
            row["Gene_symbol"] = rich_gene(clean(frame.loc[idx, "Scored_gene"]), views[idx]["is_case"])
            ws.append([
                re.sub(r"\s*[\r\n]+\s*", " | ", value).strip()
                if isinstance((value := row.get(column, "")), str)
                else value
                for column in cols
            ])
            excel_row = ws.max_row
            # Rich transcript formatting: MANE bold; Scored_transcript blue.
            tx_cell = ws.cell(excel_row, cols.index("Transcript ID") + 1)
            raw = display[idx]["Transcript ID"]
            if isinstance(raw, CellRichText):
                tx_cell.value = raw
            for c in ws[excel_row]:
                c.font = Font(name=BODY_FONT, size=BODY_FONT_SIZE, bold=(c.column == 1))
                c.fill = PatternFill("solid", fgColor=BODY_FILL)
                rule_color = HPO_COLUMN_RULE if c.column == cols.index("Gene_HPO_terms") + 1 else COLUMN_RULE
                c.border = Border(
                    left=Side(style="thin", color=COLUMN_RULE),
                    right=Side(style="thin", color=rule_color),
                    top=Side(style="thin", color=GRID),
                    bottom=Side(style="thin", color=GRID),
                )
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False, readingOrder=1)
            score_cell = ws.cell(excel_row, cols.index("Final_score") + 1)
            score_cell.font = Font(name=BODY_FONT, size=12, bold=True)
            if is_mitochondrial_variant(frame.loc[idx]):
                score_cell.comment = Comment(
                    "Partial mitochondrial evaluation using available evidence and unchanged scoring weights. MT-specific pathogenicity is unavailable. Rarity is unavailable and uses the existing missing default: 50/100, contributing 5 final-score points.",
                    "Workflow",
                )
                ws.cell(excel_row, cols.index("Rarity_score") + 1).comment = Comment(
                    "Unavailable mitochondrial rarity evidence; 50/100 is the existing missing-evidence default, not an observed rarity estimate. MT frequencies are retained for review and are not interpreted as nuclear carrier counts.",
                    "Workflow",
                )
            # Final score and position fields remain numeric/sortable.
            for name in ["Pos", "gnomAD_AF", "CADD", "LOEUF", "Pathogenicity_evidence_coverage",
                         "Variant_pathogenicity_score", "Inheritance_score", "Gene_disease_evidence_score",
                         "Case_relevance_score", "Rarity_score", "Final_score"]:
                c = ws.cell(excel_row, cols.index(name) + 1)
                if isinstance(c.value, (int, float)):
                    if name == "Pos":
                        c.number_format = "0"
                    elif name == "Pathogenicity_evidence_coverage":
                        c.number_format = "0" if isinstance(c.value, int) else "0.##"
                    elif name == "gnomAD_AF":
                        c.number_format = "0.00E+00;-0.00E+00;0"
                    elif name.endswith("score") or name == "Final_score":
                        c.number_format = "0.00"
                    else:
                        c.number_format = "0.###"
            ws.row_dimensions[excel_row].height = 22
        # Header style, widths, filters and frozen identifier columns.
        for cell in ws[1]:
            cell.value = display_header(str(cell.value))
            cell.font = Font(name=BODY_FONT, size=HEADER_FONT_SIZE, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=DARK)
            rule_color = HPO_COLUMN_RULE if cell.column == cols.index("Gene_HPO_terms") + 1 else COLUMN_RULE
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False, readingOrder=1)
            cell.border = Border(
                left=Side(style="thin", color=COLUMN_RULE),
                right=Side(style="thin", color=rule_color),
                bottom=Side(style="thin", color=GRID),
            )
        ws.row_dimensions[1].height = 28
        for col_num, name in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(col_num)].width = WIDTHS.get(name, 18)
        hidden_columns = []
        for col_num, name in enumerate(cols, 1):
            if col_num <= 5:
                continue
            if all(is_missing(ws.cell(row_num, col_num).value) for row_num in range(2, ws.max_row + 1)):
                ws.column_dimensions[get_column_letter(col_num)].hidden = True
                hidden_columns.append(name)
        if hidden_columns:
            print(f"[Stage 7] {sheet_name}: hiding empty columns: {', '.join(hidden_columns)}.", flush=True)
        ws.freeze_panes = "F2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = True
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.tabColor = BLUE if sheet_name in {"Top_candidates", "P_LP"} else "70AD47"

    output.parent.mkdir(parents=True, exist_ok=True)
    print(f"[Stage 7] Saving workbook to {output}...", flush=True)
    if tabs["ALL"] != list(frame.index) or wb["ALL"].max_row != len(frame) + 1:
        raise ValueError("Stage 7 complete review lost or duplicated variants")
    wb.save(output)
    record_qc("stage7_report", len(frame), len(frame), len(frame), len(frame), "report", tabs={name: len(indices) for name, indices in tabs.items()})
    print("[Stage 7] Workbook saved.", flush=True)


def main() -> None:
    # Snakemake passes this object into the script's global namespace.
    snakemake_obj = globals().get("snakemake")
    if snakemake_obj is None:
        raise RuntimeError("This script is intended to run from its Snakemake rule.")

    print("[Stage 7] Starting final report.", flush=True)
    print("[Stage 7] Loading Stage 6 table and case context...", flush=True)
    frame, case_context = load_inputs(
        Path(str(snakemake_obj.input.variants)),
        Path(str(snakemake_obj.input.case_context)),
    )
    print(f"[Stage 7] Loaded {len(frame):,} Stage 6 rows.", flush=True)
    write_workbook(
        frame,
        case_context,
        Path(str(snakemake_obj.output.workbook)),
    )
    print(f"[Stage 7] Complete: {snakemake_obj.output.workbook}", flush=True)


if __name__ == "__main__":
    main()
